# -*- coding: utf-8 -*-
# 论文格式改写器 —— 把 Word 论文按格式规范重新排版
# Copyright (C) 2026 林弘毅
#
# 本程序是自由软件：你可以再分发和/或修改它，遵循自由软件基金会发布的
# GNU 通用公共许可证条款，要么版本 3，要么（按你的选择）任何更新的版本。
#
# 本程序的分发是希望它有用，但不提供任何担保，甚至不提供适销性
# 或特定用途适用性的默示担保。详见 GNU 通用公共许可证。
#
# 你应该已经随本程序收到了一份 GNU 通用公共许可证的副本。
# 如果没有，见 <https://www.gnu.org/licenses/>。
"""改写结果验证：文字完整性、资产保留、格式是否真的生效"""
import sys, zipfile, collections
from lxml import etree
from docx import Document


# 保证中文输出在任何终端都不崩。
# Windows 上 stdout 默认编码可能是 cp1252，直接 print 中文会抛
# UnicodeEncodeError 并让脚本崩掉（GitHub Actions 的 windows runner 实测踩过）。
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def snapshot(path, label):
    doc = Document(path)
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        dx = z.read("word/document.xml").decode("utf-8", "ignore")
    text_all = "\n".join(p.text for p in doc.paragraphs)
    return {
        "label": label,
        "path": path,
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "sections": len(doc.sections),
        "chars": len(text_all),
        "text": text_all,
        "media": sorted(n for n in names if n.startswith("word/media/")),
        "fonts": len([n for n in names if n.endswith(".odttf")]),
        "headers": len([n for n in names if "header" in n]),
        "footers": len([n for n in names if "footer" in n]),
        "toc_field": "instrText" in dx and "TOC" in dx,
        "drawings": dx.count("<w:drawing"),
        "numbering": dx.count("<w:numPr>"),
        "bookmarks": dx.count("<w:bookmarkStart"),
        "doc": doc,
    }


def fmt_of_para(p):
    rpr = p.runs[0]._element.rPr if p.runs else None
    rf = rpr.find(W + "rFonts") if rpr is not None else None
    sz = rpr.find(W + "sz") if rpr is not None else None
    return {
        "ea": (rf.get(W + "eastAsia") if rf is not None else None),
        "ascii": (rf.get(W + "ascii") if rf is not None else None),
        "pt": (int(sz.get(W + "val")) / 2.0 if sz is not None and sz.get(W + "val") else None),
        "bold": (rpr is not None and rpr.find(W + "b") is not None),
    }


src, dst = sys.argv[1], sys.argv[2]
A, B = snapshot(src, "原件"), snapshot(dst, "改写后")

print("=" * 78)
print("一、完整性检查（改写绝不能丢东西）")
print("=" * 78)
checks = [
    ("段落数", "paragraphs"), ("表格数", "tables"), ("分节数", "sections"),
    ("正文字符数", "chars"), ("媒体/图片数", None), ("嵌入字体", "fonts"),
    ("页眉部件", "headers"), ("页脚部件", "footers"),
    ("drawing 元素", "drawings"), ("编号 numPr", "numbering"), ("书签", "bookmarks"),
]
ok = True
for name, key in checks:
    if key is None:
        a, b = len(A["media"]), len(B["media"])
    else:
        a, b = A[key], B[key]
    flag = "OK " if a == b else "!! "
    if a != b:
        ok = False
    print("  %s%-14s 原件 %-8s 改写后 %-8s" % (flag, name, a, b))

print("\n  TOC 目录域:  原件 %s  改写后 %s" % (A["toc_field"], B["toc_field"]))
if A["toc_field"] != B["toc_field"]:
    ok = False

# 文字是否一字不差
same_text = A["text"] == B["text"]
print("\n  文字内容完全一致: %s" % ("是" if same_text else "否 !!"))
if not same_text:
    ok = False
    la, lb = A["text"].split("\n"), B["text"].split("\n")
    for i, (x, y) in enumerate(zip(la, lb)):
        if x != y:
            print("     首个差异在第 %d 段:" % i)
            print("       原件  : %r" % x[:60])
            print("       改写后: %r" % y[:60])
            break

# 图片二进制是否逐字节一致
with zipfile.ZipFile(src) as za, zipfile.ZipFile(dst) as zb:
    common = set(A["media"]) & set(B["media"])
    diff = [m for m in common if za.read(m) != zb.read(m)]
    print("  图片二进制一致: %s (%d 个, 差异 %d 个)" % (
        "是" if not diff else "否 !!", len(common), len(diff)))
    if diff:
        ok = False

print("\n" + "=" * 78)
print("二、格式是否真的生效（抽检各区域代表段落）")
print("=" * 78)
doc_a, doc_b = A["doc"], B["doc"]
# 按原件的区域判定挑代表段落
sys.path.insert(0, ".")
from region_detector import detect_regions
paras, summary = detect_regions(src)
seen = set()
for info in paras:
    r = info["region"]
    if r in seen or r in ("blank", "cover_field", "toc"):
        continue
    seen.add(r)
    fa, fb = fmt_of_para(doc_a.paragraphs[info["idx"]]), fmt_of_para(doc_b.paragraphs[info["idx"]])
    print("\n  [%s] #%d %s" % (r, info["idx"], info["text"][:34]))
    print("     改写前: %spt %s 加粗=%s" % (fa["pt"], fa["ea"], fa["bold"]))
    print("     改写后: %spt %s 加粗=%s" % (fb["pt"], fb["ea"], fb["bold"]))

print("\n" + "=" * 78)
print("结论: %s" % ("完整性通过，可以检查成稿" if ok else "!! 存在完整性问题，需修复"))
print("=" * 78)
