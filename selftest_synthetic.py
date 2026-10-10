# -*- coding: utf-8 -*-
# 论文格式改写器 —— 把 Word 论文按格式规范重新排版
# Copyright (C) 2026 林弘毅
#
# 本程序是自由软件：你可以再分发和/或修改它，遵循自由软件基金会发布的
# GNU 通用公共许可证条款，要么版本 3，要么（按你的选择）任何更新的版本。
#
# 本程序的分发是希望它有用，但不提供任何担保，甚至不提供适销性
# 或特定用途适用性的默示担保。详见 GNU 通用公共许可证。
#
# 你应该已经随本程序收到了一份 GNU 通用公共许可证的副本。
# 如果没有，见 <https://www.gnu.org/licenses/>。
"""
自检：造一份假想论文，跑完整流程，自动核对每个环节是否正确
==========================================================
覆盖的边角情况：
  * 纯文本段落（无任何格式）—— 实测中最容易整片漏判
  * 编号标题（一、/（一）/1、）与无编号结构词（结语）
  * 参考文献条目里的 [1]（必须保持正常大小，不能上标）
  * 正文中 [1] 与正文同处一个 run（必须能切开）
  * 段首条目式 [1] 作者...（是条目序号，不是引用）
  * 全角括号［1］、多引用 [1,2]、区间 [3-5]
  * 注释序号 ①②（规范要求上标）
  * 标题里的 （一） 不能被当成引用上标
  * 图表题注里的「图1」不能误判

用法: python selftest_synthetic.py
"""
from __future__ import annotations

import os
import re
import sys
import tempfile
import traceback


# 保证中文输出在任何终端都不崩。
# Windows 上 stdout 默认编码可能是 cp1252，直接 print 中文会抛
# UnicodeEncodeError 并让脚本崩掉（GitHub Actions 的 windows runner 实测踩过）。
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


# ============================================================
# 辅助：往段落里加一个带显式格式的 run
# ============================================================
def add_run(p, text, ea=None, pt=None, bold=False):
    r = p.add_run(text)
    if ea is None and pt is None and not bold:
        return r                      # 故意不设格式 => 纯文本段
    rpr = r._element.get_or_add_rPr()
    if ea:
        rf = rpr.makeelement(qn("w:rFonts"), {})
        rf.set(qn("w:eastAsia"), ea)
        rf.set(qn("w:ascii"), "Times New Roman")
        rpr.insert(0, rf)
    if pt:
        sz = rpr.makeelement(qn("w:sz"), {})
        sz.set(qn("w:val"), str(int(pt * 2)))
        rpr.append(sz)
    if bold:
        rpr.append(rpr.makeelement(qn("w:b"), {}))
    return r


# ============================================================
# 造一份假想论文
# ============================================================
def build_synthetic(path: str) -> tuple[list[dict], dict]:
    """返回 (期望检查点列表, 关键段落索引字典)。"""
    doc = Document()
    expects = []
    key = {}

    def para():
        return doc.add_paragraph()

    def mark(name):
        key[name] = len(doc.paragraphs) - 1

    # ---------- 1) 纯文本区：完全没有任何格式 ----------
    p = para(); add_run(p, "数字治理视域下基层负担的生成机制研究")
    expects.append({"idx": len(doc.paragraphs) - 1, "region": "title",
                    "why": "纯文本首段，应判为大标题"})

    p = para(); add_run(p, "本文以某镇为案例，考察数字技术如何重塑基层治理。")
    expects.append({"idx": len(doc.paragraphs) - 1, "region": "body", "why": "纯文本正文"})

    p = para(); add_run(p, "一、问题的提出")
    expects.append({"idx": len(doc.paragraphs) - 1, "region": "h1",
                    "why": "纯文本一级标题（无格式），必须靠编号识别"})

    p = para(); add_run(p, "既有研究主要从制度与技术两条路径展开。")
    expects.append({"idx": len(doc.paragraphs) - 1, "region": "body", "why": "纯文本正文"})

    p = para(); add_run(p, "（一）制度路径")
    expects.append({"idx": len(doc.paragraphs) - 1, "region": "h2",
                    "why": "纯文本二级标题"})

    p = para(); add_run(p, "制度路径强调压力型体制下的层层加码。")
    expects.append({"idx": len(doc.paragraphs) - 1, "region": "body", "why": "纯文本正文"})

    p = para(); add_run(p, "（二）技术路径")
    expects.append({"idx": len(doc.paragraphs) - 1, "region": "h2",
                    "why": "纯文本二级标题"})

    # ---------- 2) 带格式区 ----------
    p = para(); add_run(p, "二、案例分析", ea="黑体", pt=14)
    expects.append({"idx": len(doc.paragraphs) - 1, "region": "h1",
                    "why": "黑体14pt，格式信号"})

    # 正文里含引用标记，且标记与正文在同一个 run
    p = para(); add_run(p, "已有研究指出该问题具有普遍性[1]，但在机制上仍有争议[2]。",
                        ea="宋体", pt=12)
    idx_mixed = len(doc.paragraphs) - 1
    expects.append({"idx": idx_mixed, "region": "body",
                    "why": "正文（含 [1][2]，与正文同 run，需切开）"})
    expects.append({"idx": idx_mixed, "superscript": ["[1]", "[2]"],
                    "why": "run 内引用必须被切开并上标"})

    # 多引用与区间
    p = para(); add_run(p, "相关研究见[1,2]与[3-5]。", ea="宋体", pt=12)
    idx_multi = len(doc.paragraphs) - 1
    expects.append({"idx": idx_multi, "region": "body", "why": "正文"})
    expects.append({"idx": idx_multi, "superscript": ["[1,2]", "[3-5]"],
                    "why": "多引用与区间都要上标"})

    # 全角括号
    p = para(); add_run(p, "另有学者持相反意见［6］。", ea="宋体", pt=12)
    idx_full = len(doc.paragraphs) - 1
    expects.append({"idx": idx_full, "superscript": ["［6］"],
                    "why": "全角括号引用也要上标"})

    # 注释序号
    p = para(); add_run(p, "该提法最早见于马克思的手稿①，后经恩格斯整理②。", ea="宋体", pt=12)
    idx_note = len(doc.paragraphs) - 1
    expects.append({"idx": idx_note, "superscript": ["①", "②"],
                    "why": "注释序号规范要求上标"})

    # 图表题注
    p = para(); add_run(p, "图1：分析框架示意", ea="宋体", pt=10.5)
    expects.append({"idx": len(doc.paragraphs) - 1, "region": "caption",
                    "why": "图表题注"})

    # ---------- 3) 标题必须不被上标 ----------
    p = para(); add_run(p, "三、结论与讨论", ea="黑体", pt=14)
    idx_h = len(doc.paragraphs) - 1
    mark("h1_formatted")
    expects.append({"idx": idx_h, "region": "h1", "why": "带格式一级标题"})
    expects.append({"idx": idx_h, "superscript": [],
                    "why": "标题里没有引用标记，不应有任何上标"})

    # 无编号结构词
    p = para(); add_run(p, "结语")
    expects.append({"idx": len(doc.paragraphs) - 1, "region": "h1",
                    "why": "无编号结构词，纯文本也要能识别"})

    p = para(); add_run(p, "本研究认为，数字技术的影响是双重的。", ea="宋体", pt=12)
    expects.append({"idx": len(doc.paragraphs) - 1, "region": "body", "why": "正文"})

    # ---------- 5) 刁钻边角（必须在参考文献区之前，否则会被判为条目）----------
    # 5a 引用后紧跟句号
    p = para(); add_run(p, "这一判断在后续研究中被反复验证[7]。", ea="宋体", pt=12)
    mark("cite_before_period")
    expects.append({"idx": key["cite_before_period"], "region": "body", "why": "正文"})
    expects.append({"idx": key["cite_before_period"], "superscript": ["[7]"],
                    "why": "引用后紧跟句号"})

    # 5b 同一段落里既有引用又有注释序号
    p = para(); add_run(p, "参见前文①的讨论[8]，此处不再展开。", ea="宋体", pt=12)
    mark("mixed_note_cite")
    expects.append({"idx": key["mixed_note_cite"], "superscript": ["①", "[8]"],
                    "why": "注释与引用混排"})

    # 5c 年份区间不能被当成引用（数字位数与引用不同，可区分）
    p = para(); add_run(p, "统计区间为[2010-2015]年。", ea="宋体", pt=12)
    mark("year_range")
    expects.append({"idx": key["year_range"], "region": "body", "why": "正文"})
    expects.append({"idx": key["year_range"], "superscript": [],
                    "why": "★ 年份区间 [2010-2015] 不是引用，不能上标"})

    # 5d 【摘要】标注 + 内容同段
    p = para(); add_run(p, "【摘要】：本文以某镇为例，考察数字治理下的基层负担[9]。")
    mark("abstract_inline")
    expects.append({"idx": key["abstract_inline"], "region": "abstract",
                    "why": "标注与内容同段，应整段判为摘要"})
    expects.append({"idx": key["abstract_inline"], "superscript": ["[9]"],
                    "why": "摘要里的引用也要上标"})

    # 5e 关键词行
    p = para(); add_run(p, "【关键词】数字治理；基层负担；形式主义")
    mark("keywords_inline")
    expects.append({"idx": key["keywords_inline"], "region": "keywords",
                    "why": "关键词标注与内容同段"})

    # 5f 英文摘要标题与内容同段
    p = para(); add_run(p, "Abstract: This paper examines digital governance.")
    mark("abs_en")
    expects.append({"idx": key["abs_en"], "region": "abstract",
                    "why": "英文摘要「Abstract: 内容」同段，应判为摘要内容"})

    # ---------- 6) 参考文献区（必须是文档最后，规范顺序如此）----------
    p = para(); add_run(p, "参考文献", ea="黑体", pt=14)
    idx_rh = len(doc.paragraphs) - 1
    expects.append({"idx": idx_rh, "region": "ref_head", "why": "参考文献标题"})

    p = para(); add_run(p, "[1] 汪信砚. 马克思思想的形成过程[J]. 哲学研究, 2010(02):3-10.",
                        ea="宋体", pt=9)
    idx_r1 = len(doc.paragraphs) - 1
    mark("ref_item_1")
    expects.append({"idx": idx_r1, "region": "ref_item",
                    "why": "参考文献条目"})
    expects.append({"idx": idx_r1, "superscript": [],
                    "why": "★ 条目里的 [1] 是序号，规范要求左顶格，绝不能上标"})

    p = para(); add_run(p, "[2] 佟林杰, 张明欣. 基层数字形式主义研究[J]. 学术交流, 2022(5):1-9.",
                        ea="宋体", pt=9)
    idx_r2 = len(doc.paragraphs) - 1
    mark("ref_item_2")
    expects.append({"idx": idx_r2, "superscript": [],
                    "why": "★ 同上，条目序号不上标"})

    doc.save(path)
    return expects, key


# ============================================================
# 检查器
# ============================================================
def read_run_state(path: str):
    """读回文档：每个段落的区域无关信息 + 每个 run 的文本与是否上标。"""
    doc = Document(path)
    out = []
    for p in doc.paragraphs:
        runs = []
        for r in p.runs:
            rpr = r._element.rPr
            va = rpr.find(qn("w:vertAlign")) if rpr is not None else None
            sup = va is not None and va.get(qn("w:val")) == "superscript"
            runs.append({"text": r.text, "sup": sup})
        out.append({"text": p.text, "runs": runs})
    return doc, out


def main():
    tmp = tempfile.mkdtemp(prefix="synth_")
    src = os.path.join(tmp, "假想论文.docx")
    out = os.path.join(tmp, "假想论文_已排版.docx")

    print("=" * 78)
    print("步骤 1：造一份假想论文（含纯文本段、引用、注释、参考文献）")
    print("=" * 78)
    expects, key = build_synthetic(src)
    doc_src = Document(src)
    print("  段落数: %d，检查点: %d 条" % (len(doc_src.paragraphs), len(expects)))

    print("\n" + "=" * 78)
    print("步骤 2：跑区域识别")
    print("=" * 78)
    from region_detector import detect_regions
    paras, summary = detect_regions(src)
    print("  识别结果:", {k: v for k, v in summary["counts"].items() if v})

    print("\n" + "=" * 78)
    print("步骤 3：跑完整改写（含上标）")
    print("=" * 78)
    from rewriter import rewrite
    res = rewrite(src, out, renumber_headings=False,
                  superscript_citations=True, superscript_notes=True)
    print("  改写段落: %d" % res["changed_count"])
    sup = res.get("superscript") or {}
    print("  上标标记: %d 处（涉及 %d 段）" % (
        sup.get("markers_superscripted", 0), sup.get("paragraphs_touched", 0)))
    print("  跳过参考文献条目: %d 段" % sup.get("skipped_ref_items", 0))
    print("  跳过条目式开头: %d 段" % sup.get("skipped_entry_like", 0))

    print("\n" + "=" * 78)
    print("步骤 4：逐条核对检查点")
    print("=" * 78)
    region_of = {p["idx"]: p["region"] for p in paras}
    doc_out, state = read_run_state(out)

    passed = failed = 0
    failures = []

    # 4a 区域判定
    for e in expects:
        if "region" not in e:
            continue
        got = region_of.get(e["idx"])
        if got == e["region"]:
            passed += 1
        else:
            failed += 1
            failures.append("区域：段落 #%d 期望 %s，实际 %s（%s）  文本=%r" % (
                e["idx"], e["region"], got, e["why"], state[e["idx"]]["text"][:40]))

    # 4b 上标
    known_limits = []
    for e in expects:
        if "superscript" not in e:
            continue
        want = e["superscript"]
        got = [r["text"] for r in state[e["idx"]]["runs"] if r["sup"]]
        if sorted(got) == sorted(want):
            passed += 1
        elif e.get("known_limitation"):
            known_limits.append("已知局限：段落 #%d 期望 %s，实际 %s（%s）" % (
                e["idx"], want, got, e["why"]))
        else:
            failed += 1
            failures.append("上标：段落 #%d 期望 %s，实际 %s（%s）  文本=%r" % (
                e["idx"], want, got, e["why"], state[e["idx"]]["text"][:40]))

    # 4c 文字不能丢
    src_texts = [p.text for p in doc_src.paragraphs]
    out_texts = [p.text for p in doc_out.paragraphs]
    if src_texts == out_texts:
        passed += 1
        print("  [OK] 文字完整性：%d 段全部一字不差" % len(src_texts))
    else:
        failed += 1
        for i, (a, b) in enumerate(zip(src_texts, out_texts)):
            if a != b:
                failures.append("文字被改动：段落 #%d\n      原: %r\n      新: %r" % (i, a[:60], b[:60]))
                break
        failures.append("文字完整性检查失败")

    # 4d 参考文献条目里的 [n] 必须仍然不是上标（最关键的一条）
    for i in (key["ref_item_1"], key["ref_item_2"]):
        sup_runs = [r for r in state[i]["runs"] if r["sup"]]
        if not sup_runs:
            passed += 1
        else:
            failed += 1
            failures.append("严重：参考文献条目 #%d 的序号被误标为上标 %s" % (
                i, [r["text"] for r in sup_runs]))

    # 4e 标题里不应出现任何上标
    for i in (key["h1_formatted"],):
        if not [r for r in state[i]["runs"] if r["sup"]]:
            passed += 1
        else:
            failed += 1
            failures.append("标题 #%d 被误加上标" % i)

    print("\n" + "=" * 78)
    print("结果：通过 %d 项，失败 %d 项" % (passed, failed))
    print("=" * 78)
    if failures:
        print("\n失败明细：")
        for f in failures:
            print("  ✗ %s" % f)
    if known_limits:
        print("\n已知局限（不计入失败，需人工判断）：")
        for k in known_limits:
            print("  ! %s" % k)
    if not failures:
        print("  全部通过")

    print("\n产物: %s" % out)
    return 0 if failed == 0 else 1


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception:
        traceback.print_exc()
        sys.exit(2)
