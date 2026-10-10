# -*- coding: utf-8 -*-
"""
验收检验：模拟全新用户从拿到程序到生成成稿的完整流程
====================================================
不是单元测试，而是端到端验收：
  1. 环境依赖是否齐全
  2. 造一份贴近真实难度的论文（含各种边角情况）
  3. 走完「识别 → 改写」全流程
  4. 逐项核对输出：字体、字号、编号、上标、文字完整性
  5. 报告通过/失败

用法: python 验收检验.py
"""
from __future__ import annotations

import os
import re
import shutil
import sys
import tempfile
import traceback


# 保证中文输出在任何终端都不崩。
# Windows 上 stdout 默认编码可能是 cp1252，直接 print 中文会抛
# UnicodeEncodeError 并让脚本崩掉（GitHub Actions 的 windows runner 实测踩过）。
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

PASS = 0
FAIL = 0
FAILURES: list[str] = []
SECTION = ""


def sec(title: str):
    global SECTION
    SECTION = title
    print("\n" + "=" * 74)
    print(title)
    print("=" * 74)


def ok(name: str, cond: bool, detail: str = ""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print("  [通过] %s" % name)
    else:
        FAIL += 1
        FAILURES.append("[%s] %s  %s" % (SECTION, name, detail))
        print("  [失败] %s   %s" % (name, detail))


def add_run(p, text, ea=None, pt=None, bold=False, sup=False):
    r = p.add_run(text)
    if ea is None and pt is None and not bold and not sup:
        return r
    rpr = r._element.get_or_add_rPr()
    if ea:
        rf = rpr.makeelement(qn("w:rFonts"), {})
        rf.set(qn("w:eastAsia"), ea)
        rf.set(qn("w:ascii"), "Times New Roman")
        rpr.insert(0, rf)
    if pt:
        sz = rpr.makeelement(qn("w:sz"), {})
        sz.set(qn("w:val"), str(int(pt * 2)))
        rpr.append(sz)
    if bold:
        rpr.append(rpr.makeelement(qn("w:b"), {}))
    if sup:
        va = rpr.makeelement(qn("w:vertAlign"), {})
        va.set(qn("w:val"), "superscript")
        rpr.append(va)
    return r


def build_paper(path: str) -> dict:
    """造一份覆盖边角情况的论文，返回关键段落索引。"""
    doc = Document()
    k = {}

    def par():
        return doc.add_paragraph()

    def mark(name):
        k[name] = len(doc.paragraphs) - 1

    # 纯文本区（无任何格式）
    p = par(); add_run(p, "数字治理视域下基层负担的生成机制研究"); mark("title")
    p = par(); add_run(p, "一、问题的提出"); mark("h1_plain")
    p = par(); add_run(p, "既有研究主要从制度与技术两条路径展开。"); mark("body_plain")

    # 带格式区
    p = par(); add_run(p, "二、案例分析", ea="黑体", pt=14); mark("h1_fmt")
    p = par(); add_run(p, "（一）制度路径", ea="黑体", pt=14); mark("h2_fmt")
    p = par(); add_run(p, "该问题具有普遍性[1]，但机制上仍有争议[2]。",
                       ea="宋体", pt=12); mark("body_cite")
    p = par(); add_run(p, "注释见①与②。", ea="宋体", pt=12); mark("body_note")
    p = par(); add_run(p, "图1：分析框架", ea="宋体", pt=10.5); mark("caption")
    p = par(); add_run(p, "结语"); mark("h1_word")

    # 摘要标注与内容同段
    p = par(); add_run(p, "【摘要】：本文以某镇为例考察数字治理下的基层负担[3]。")
    mark("abstract")

    # 参考文献区（必须在最后）
    p = par(); add_run(p, "参考文献", ea="黑体", pt=14); mark("ref_head")
    p = par(); add_run(p, "[1] 汪信砚. 马克思思想的形成过程[J]. 哲学研究, 2010(02):3-10.",
                       ea="宋体", pt=9); mark("ref1")
    p = par(); add_run(p, "[2] 佟林杰, 张明欣. 基层数字形式主义研究[J]. 学术交流, 2022(5):1-9.",
                       ea="宋体", pt=9); mark("ref2")

    doc.save(path)
    return k


def para_font(p):
    """取段落首个非空 run 的字体信息。"""
    for r in p.runs:
        if not (r.text and r.text.strip()):
            continue
        rpr = r._element.rPr
        rf = rpr.find(qn("w:rFonts")) if rpr is not None else None
        sz = rpr.find(qn("w:sz")) if rpr is not None else None
        return {
            "ea": rf.get(qn("w:eastAsia")) if rf is not None else None,
            "pt": (int(sz.get(qn("w:val"))) / 2.0) if sz is not None and sz.get(qn("w:val")) else None,
            "bold": (rpr is not None and rpr.find(qn("w:b")) is not None),
        }
    return {"ea": None, "pt": None, "bold": False}


def superscripts(p):
    out = []
    for r in p.runs:
        rpr = r._element.rPr
        va = rpr.find(qn("w:vertAlign")) if rpr is not None else None
        if va is not None and va.get(qn("w:val")) == "superscript":
            out.append(r.text)
    return out


def main() -> int:
    print("=" * 74)
    print("验收检验 —— 论文格式改写器")
    print("=" * 74)
    print("Python: %s" % sys.version.split()[0])
    print("平台  : %s" % sys.platform)

    # ---------- 1. 依赖 ----------
    sec("1. 环境依赖")
    for mod, name in (("docx", "python-docx"), ("lxml.etree", "lxml")):
        try:
            __import__(mod)
            ok(name, True)
        except ImportError as e:
            ok(name, False, str(e))
    # 注意：tkinter 的导入本身在无显示环境也能成功，所以这里只检查「能否导入」。
    # 真正建窗口的检查在第 10 节，那里会先探测显示环境。
    try:
        import tkinter  # noqa
        ok("tkinter（图形界面库可导入）", True)
    except Exception as e:
        ok("tkinter（图形界面库可导入）", False, str(e))

    # ---------- 2. 造论文 ----------
    sec("2. 构造测试论文（覆盖边角情况）")
    tmp = tempfile.mkdtemp(prefix="acceptance_")
    src = os.path.join(tmp, "测试论文.docx")
    out = os.path.join(tmp, "测试论文_已排版.docx")
    key = build_paper(src)
    doc_src = Document(src)
    ok("论文生成（%d 段）" % len(doc_src.paragraphs), True)

    # ---------- 3. 区域识别 ----------
    sec("3. 区域识别")
    from region_detector import detect_regions
    paras, summary = detect_regions(src)
    regions = {p["idx"]: p["region"] for p in paras}
    counts = summary["counts"]
    print("  识别分布: %s" % {k2: v for k2, v in counts.items() if v})

    expect_regions = [
        ("title", "大标题（纯文本）"),
        ("h1_plain", "一级标题（纯文本，靠编号）"),
        ("body_plain", "正文（纯文本）"),
        ("h1_fmt", "一级标题（黑体14pt）"),
        ("h2_fmt", "二级标题（黑体14pt）"),
        ("body_cite", "正文（含引用）"),
        ("caption", "图表题注"),
        ("h1_word", "结构词「结语」"),
        ("abstract", "摘要标注与内容同段"),
        ("ref_head", "参考文献标题"),
        ("ref1", "参考文献条目"),
    ]
    for kk, desc in expect_regions:
        exp = {"title": "title", "h1_plain": "h1", "body_plain": "body",
               "h1_fmt": "h1", "h2_fmt": "h2", "body_cite": "body",
               "caption": "caption", "h1_word": "h1", "abstract": "abstract",
               "ref_head": "ref_head", "ref1": "ref_item"}[kk]
        got = regions.get(key[kk])
        ok("%s -> %s" % (desc, exp), got == exp, "实际 %s" % got)

    # ---------- 4. 改写 ----------
    sec("4. 执行改写")
    from rewriter import rewrite
    src_texts = [p.text for p in doc_src.paragraphs]
    res = rewrite(src, out, renumber_headings=True)
    print("  改写段落: %d" % res["changed_count"])
    print("  编号处理: 去掉自动编号 %d，重排 %d" % (
        (res.get("numbering") or {}).get("stripped_auto_numbering", 0),
        (res.get("numbering") or {}).get("renumbered", 0)))
    print("  上标处理: %d 处，跳过条目 %d 段" % (
        (res.get("superscript") or {}).get("markers_superscripted", 0),
        (res.get("superscript") or {}).get("skipped_ref_items", 0)))
    ok("生成了输出文件", os.path.exists(out))

    doc_out = Document(out)
    out_texts = [p.text for p in doc_out.paragraphs]

    # ---------- 5. 字体字号 ----------
    sec("5. 字体字号是否生效")
    checks = [
        ("h1_fmt", "黑体", 14.0, False, "一级标题"),
        ("h2_fmt", "黑体", 14.0, False, "二级标题"),
        ("body_cite", "宋体", 12.0, False, "正文"),
        ("caption", "宋体", 10.5, False, "图表题注"),
        ("ref1", "宋体", 9.0, False, "参考文献条目"),
    ]
    for kk, ea, pt, bold, desc in checks:
        f = para_font(doc_out.paragraphs[key[kk]])
        good = (f["ea"] == ea and f["pt"] == pt and f["bold"] == bold)
        ok("%s = %s %.1fpt 粗=%s" % (desc, ea, pt, bold), good,
           "实际 %s %s pt=%s" % (f["ea"], f["pt"], f["bold"]))

    # ---------- 6. 编号重排 ----------
    sec("6. 标题编号重排")
    t_h1 = doc_out.paragraphs[key["h1_plain"]].text
    t_h1f = doc_out.paragraphs[key["h1_fmt"]].text
    t_h2f = doc_out.paragraphs[key["h2_fmt"]].text
    t_word = doc_out.paragraphs[key["h1_word"]].text
    print("  一级标题: %r / %r" % (t_h1, t_h1f))
    print("  二级标题: %r" % t_h2f)
    print("  结构词  : %r" % t_word)
    ok("纯文本一级标题被编号", re.match(r"^[一二三四五六七八九十]+、", t_h1) is not None,
       "实际 %r" % t_h1)
    ok("带格式一级标题被编号", re.match(r"^[一二三四五六七八九十]+、", t_h1f) is not None,
       "实际 %r" % t_h1f)
    ok("二级标题为（一）式", re.match(r"^（[一二三四五六七八九十]+）", t_h2f) is not None,
       "实际 %r" % t_h2f)
    ok("结构词标题也被编号", re.match(r"^[一二三四五六七八九十]+、", t_word) is not None,
       "实际 %r" % t_word)
    # 编号连续性
    nums = []
    for i in (key["h1_plain"], key["h1_fmt"], key["h1_word"]):
        m = re.match(r"^([一二三四五六七八九十]+)、", doc_out.paragraphs[i].text)
        if m:
            nums.append(m.group(1))
    cn = "一二三四五六七八九十"
    contiguous = all(cn.find(b) == cn.find(a) + 1 for a, b in zip(nums, nums[1:]))
    ok("一级标题编号连续 %s" % nums, contiguous and len(nums) == 3)

    # ---------- 7. 上标 ----------
    sec("7. 上标处理")
    sup_cite = superscripts(doc_out.paragraphs[key["body_cite"]])
    sup_note = superscripts(doc_out.paragraphs[key["body_note"]])
    sup_ref1 = superscripts(doc_out.paragraphs[key["ref1"]])
    sup_abs = superscripts(doc_out.paragraphs[key["abstract"]])
    print("  正文引用上标: %s" % sup_cite)
    print("  注释序号上标: %s" % sup_note)
    print("  摘要引用上标: %s" % sup_abs)
    print("  参考文献条目上标: %s（应为空）" % sup_ref1)
    ok("正文引用 [1][2] 已上标", "[1]" in sup_cite and "[2]" in sup_cite, str(sup_cite))
    ok("注释序号 ① 已上标", "①" in sup_note, str(sup_note))
    ok("摘要里的引用也上标", "[3]" in sup_abs, str(sup_abs))
    ok("★ 参考文献条目序号未上标", sup_ref1 == [], str(sup_ref1))

    # ---------- 8. 文字完整性 ----------
    sec("8. 文字完整性")
    diff = [(i, a, b) for i, (a, b) in enumerate(zip(src_texts, out_texts)) if a != b]
    ok("段落数一致", len(src_texts) == len(out_texts),
       "%d vs %d" % (len(src_texts), len(out_texts)))
    # 允许的差异：标题前缀被补写。逐条确认差异都是编号前缀
    unexpected = []
    for i, a, b in diff:
        stripped = re.sub(r"^([一二三四五六七八九十]+、|（[一二三四五六七八九十]+）|\d+、)",
                          "", b)
        if stripped != a:
            unexpected.append((i, a[:30], b[:30]))
    ok("差异仅为补写的编号前缀（%d 处）" % len(diff), not unexpected, str(unexpected[:3]))
    body_same = src_texts[key["body_cite"]] == out_texts[key["body_cite"]]
    ok("正文文字未被改动", body_same,
       "%r vs %r" % (src_texts[key["body_cite"]][:30], out_texts[key["body_cite"]][:30]))
    ref_same = src_texts[key["ref1"]] == out_texts[key["ref1"]]
    ok("参考文献文字未被改动", ref_same)

    # ---------- 9. 原文件未被改动 ----------
    sec("9. 原文件安全性")
    doc_check = Document(src)
    same = [p.text for p in doc_check.paragraphs] == src_texts
    ok("原文件内容未变", same)
    src_font = para_font(doc_check.paragraphs[key["h1_fmt"]])
    ok("原文件格式未变（仍为黑体14pt）",
       src_font["ea"] == "黑体" and src_font["pt"] == 14.0, str(src_font))

    # ---------- 10. GUI 可构建 ----------
    sec("10. 图形界面")
    # 无图形显示的环境（CI runner、纯 SSH）建不了窗口，这属于环境限制而非程序缺陷，
    # 应当跳过而不是判为失败。踩过的坑：GitHub ubuntu runner 上报
    # TclError: no display name and no $DISPLAY environment variable，导致 CI 标红。
    try:
        from environment_probe import can_create_window
        can_gui, why = can_create_window()
    except Exception:
        can_gui, why = True, ""      # 探测模块缺失时按旧行为尝试
    if not can_gui:
        print("  [跳过] 图形界面：当前环境无显示（%s）" % why)
        print("         这不算失败——请在桌面环境或有显示的机器上验证界面")
    else:
        try:
            import tkinter as tk
            import app as appmod
            root = tk.Tk()
            root.withdraw()
            a = appmod.App(root)
            n = len(a.vars)
            root.destroy()
            ok("界面构建成功（%d 个区域）" % n, n >= 10)
        except Exception as e:
            ok("界面构建", False, repr(e)[:120])

    # ---------- 汇总 ----------
    sec("结果汇总")
    print("  通过 %d 项，失败 %d 项" % (PASS, FAIL))
    if FAILURES:
        print("\n  失败明细:")
        for f in FAILURES:
            print("    ✗ %s" % f)
    else:
        print("  全部通过 —— 程序可正常使用")
    print("\n  测试产物目录: %s" % tmp)
    return 0 if FAIL == 0 else 1


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception:
        traceback.print_exc()
        sys.exit(2)
