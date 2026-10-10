# -*- coding: utf-8 -*-
# 论文格式改写器 —— 把 Word 论文按格式规范重新排版
# Copyright (C) 2026 林弘毅
#
# 本程序是自由软件：你可以再分发和/或修改它，遵循自由软件基金会发布的
# GNU 通用公共许可证条款，要么版本 3，要么（按你的选择）任何更新的版本。
#
# 本程序的分发是希望它有用，但不提供任何担保，甚至不提供适销性
# 或特定用途适用性的默示担保。详见 GNU 通用公共许可证。
#
# 你应该已经随本程序收到了一份 GNU 通用公共许可证的副本。
# 如果没有，见 <https://www.gnu.org/licenses/>。
"""
自检：跨平台字体适配
====================
重点验证几件靠猜会出错的事：
  1. 写入 XML 的字体名**不含逗号**（曾经错误地写入字体回退链，被 Word 实测推翻）
  2. Windows 上不做任何字体替换（保持已验证行为）
  3. macOS 上把不存在的字体替换成存在的
  4. 探测失败时原样返回，不引入新错误
  5. open_path 跨平台可用
  6. macOS 界面字体选项里不出现 Windows 字体名

用法: python selftest_platform.py
"""
from __future__ import annotations

import os
import sys
import traceback

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import platform_fonts as pf

PASS = FAIL = 0
FAILURES: list[str] = []


def check(name: str, cond: bool, detail: str = ""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print("  [OK]   %s" % name)
    else:
        FAIL += 1
        FAILURES.append("%s   %s" % (name, detail))
        print("  [失败] %s   %s" % (name, detail))


def main() -> int:
    print("=" * 74)
    print("跨平台字体适配自检")
    print("=" * 74)
    print("当前平台: %s" % sys.platform)

    # ---------- 1. 核心约束：字体名不含逗号 ----------
    print("\n1. 字体名不得含逗号（写入 XML 的名字是单个字体名）")
    names = ["宋体", "黑体", "楷体", "仿宋_GB2312", "微软雅黑",
             "Times New Roman", "Arial", "Songti SC", "不存在的字体XYZ"]
    bad = []
    for n in names:
        r = pf.resolve_font(n)
        if "," in r:
            bad.append((n, r))
    check("resolve_font 输出无逗号", not bad, "违规: %s" % bad)

    # 候选表本身也不能含逗号（否则会被当整体名写进去）
    bad_alias = []
    for k, cands in pf.FONT_ALIASES.items():
        for c in cands:
            if "," in c:
                bad_alias.append((k, c))
    check("候选表无逗号", not bad_alias, "违规: %s" % bad_alias)

    # ---------- 2. Windows 上不替换 ----------
    print("\n2. Windows 上保持原样（不引入未验证的改动）")
    if sys.platform == "win32":
        unchanged = []
        for n in ("宋体", "黑体", "楷体", "Times New Roman"):
            r = pf.resolve_font(n)
            if r != n:
                unchanged.append((n, r))
        check("Windows 不替换任何字体", not unchanged, "被改了: %s" % unchanged)
    else:
        print("  (跳过：当前不是 Windows)")

    # ---------- 3. 模拟 macOS 行为 ----------
    print("\n3. 模拟 macOS：应把缺失字体替换成存在的")
    saved_plat = sys.platform
    saved_cache = None
    try:
        # 伪造「系统只有 Mac 字体」的环境
        fake_have = {"songti sc", "heiti sc", "kaiti sc",
                     "times new roman", "stfangsong", "pingfang sc"}

        def fake_installed():
            return fake_have

        sys.platform = "darwin"
        pf.installed_fonts.cache_clear()
        orig_installed = pf.installed_fonts
        pf.installed_fonts = fake_installed          # type: ignore

        cases = [
            ("宋体", "Songti SC"),
            ("黑体", "Heiti SC"),
            ("楷体", "Kaiti SC"),
            ("仿宋_GB2312", "STFangsong"),
            ("Times New Roman", "Times New Roman"),   # 存在，不该改
        ]
        for src, want in cases:
            got = pf.resolve_font(src)
            check("macOS %s -> %s" % (src, want), got == want, "实际 %r" % got)

        # 完全找不到的，应原样返回而不是报错
        got = pf.resolve_font("不存在的字体XYZ")
        check("macOS 无候选时原样返回", got == "不存在的字体XYZ", "实际 %r" % got)

        # 探测失败（返回空集合）时应原样返回
        pf.installed_fonts = lambda: set()           # type: ignore
        got = pf.resolve_font("宋体")
        check("macOS 探测失败时原样返回", got == "宋体", "实际 %r" % got)

        # ---------- 4. 界面字体选项 ----------
        print("\n4. macOS 界面字体选项不应出现 Windows 字体名")
        pf.installed_fonts = fake_installed          # type: ignore
        from rewriter import FONT_CHOICES
        opts = pf.platform_font_choices(FONT_CHOICES)
        win_only = [o for o in opts if o in ("宋体", "黑体", "楷体", "仿宋_GB2312", "微软雅黑")]
        check("macOS 选项已换成 Mac 字体", not win_only, "仍含: %s" % win_only)
        check("macOS 选项包含 Songti SC", "Songti SC" in opts, "选项: %s" % opts)

        # ---------- 5. 预设批量解析 ----------
        print("\n5. 预设批量解析")
        from rewriter import DEFAULT_PRESET
        newp, changes = pf.resolve_fonts_in_preset(DEFAULT_PRESET)
        check("批量解析返回替换记录", len(changes) > 0, "记录数 %d" % len(changes))
        all_resolved = all("," not in f.font_ea for f in newp.values())
        check("解析后字体名仍无逗号", all_resolved)
        # 原预设不能被改动
        orig_ea = {k: v.font_ea for k, v in DEFAULT_PRESET.items()}
        still = all(DEFAULT_PRESET[k].font_ea == v for k, v in orig_ea.items())
        check("原预设对象未被修改", still)
        for region, old, new in changes[:5]:
            print("       %s: %s -> %s" % (region, old, new))
    finally:
        sys.platform = saved_plat
        try:
            pf.installed_fonts = orig_installed      # type: ignore
        except Exception:
            pass
        pf.installed_fonts.cache_clear()

    # ---------- 6. open_path ----------
    print("\n6. 跨平台打开文件")
    check("open_path 存在且可调用", callable(pf.open_path))
    ok = pf.open_path(os.path.join(os.path.dirname(os.path.abspath(__file__)), "README.md"))
    # 在无图形环境里可能失败，但只要不抛异常就算通过
    check("open_path 不抛异常", True, "返回值 %s" % ok)

    # ---------- 7. 真实改写产物不含逗号字体 ----------
    print("\n7. 真实改写：产物 XML 里的字体名不含逗号")
    import tempfile
    from docx import Document as _D
    from docx.oxml.ns import qn
    from rewriter import rewrite
    tmp = tempfile.mkdtemp()
    src = os.path.join(tmp, "t.docx")
    d = _D()
    d.add_paragraph("一、引言")
    d.add_paragraph("这是正文。")
    d.save(src)
    out = os.path.join(tmp, "t_out.docx")
    rewrite(src, out)
    d2 = _D(out)
    bad_fonts = []
    for p in d2.paragraphs:
        for r in p.runs:
            rpr = r._element.rPr
            rf = rpr.find(qn("w:rFonts")) if rpr is not None else None
            if rf is None:
                continue
            for attr in ("w:eastAsia", "w:ascii", "w:hAnsi"):
                v = rf.get(qn(attr))
                if v and "," in v:
                    bad_fonts.append((p.text[:12], attr, v))
    check("产物字体名不含逗号", not bad_fonts, "违规: %s" % bad_fonts[:3])

    print("\n" + "=" * 74)
    print("结果：通过 %d 项，失败 %d 项" % (PASS, FAIL))
    print("=" * 74)
    if FAILURES:
        print("\n失败明细：")
        for f in FAILURES:
            print("  ✗ %s" % f)
    else:
        print("  全部通过")
    return 0 if FAIL == 0 else 1


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception:
        traceback.print_exc()
        sys.exit(2)
