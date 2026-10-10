# -*- coding: utf-8 -*-
# 论文格式改写器 —— 把 Word 论文按格式规范重新排版
# Copyright (C) 2026 林弘毅
#
# 本程序是自由软件：你可以再分发和/或修改它，遵循自由软件基金会发布的
# GNU 通用公共许可证条款，要么版本 3，要么（按你的选择）任何更新的版本。
#
# 本程序的分发是希望它有用，但不提供任何担保，甚至不提供适销性
# 或特定用途适用性的默示担保。详见 GNU 通用公共许可证。
#
# 你应该已经随本程序收到了一份 GNU 通用公共许可证的副本。
# 如果没有，见 <https://www.gnu.org/licenses/>。
"""
程序入口
========
正常双击：启动图形界面。
带参数 --selftest：不开窗口，直接验证打包后的依赖与核心功能是否可用。
   这个模式专为验证 exe 而设——比启动 GUI 观察窗口可靠得多。
"""
import os
import sys



# 保证中文输出在任何终端都不崩。
# Windows 上 stdout 默认编码可能是 cp1252，直接 print 中文会抛
# UnicodeEncodeError 并让脚本崩掉（GitHub Actions 的 windows runner 实测踩过）。
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass
def selftest() -> int:
    lines = []
    ok = True

    def out(s=""):
        lines.append(s)
        try:
            print(s)
        except Exception:
            pass   # --windowed 打包后没有控制台，忽略

    out("=" * 62)
    out("论文格式改写器 —— 打包自检")
    out("=" * 62)

    def check(name, fn):
        nonlocal ok
        try:
            r = fn()
            out("  [OK]   %-26s %s" % (name, r if r else ""))
        except Exception as e:
            ok = False
            out("  [失败] %-26s %r" % (name, e))

    def _python():
        return "Python %s" % sys.version.split()[0]

    def _frozen():
        return "打包模式" if getattr(sys, "frozen", False) else "脚本模式"

    def _docx():
        import docx
        return "python-docx %s" % getattr(docx, "__version__", "?")

    def _lxml():
        import lxml.etree
        return "lxml %s" % lxml.etree.__version__

    def _tk():
        import tkinter
        return "Tk %s" % tkinter.TkVersion

    def _modules():
        import region_detector, rewriter, numbering  # noqa: F401
        return "区域识别 / 改写 / 编号 三个模块"

    def _preset():
        import json
        from rewriter import DEFAULT_PRESET, preset_from_dict, preset_to_dict
        s = json.dumps(preset_to_dict(DEFAULT_PRESET), ensure_ascii=False)
        back = preset_from_dict(json.loads(s))
        assert set(back) == set(DEFAULT_PRESET)
        return "默认预设 %d 个区域，序列化往返一致" % len(DEFAULT_PRESET)

    def _writable():
        from app import user_dir, preset_path
        d = user_dir()
        assert os.path.isdir(d), "目录不存在: %s" % d
        p = preset_path()
        with open(p, "a", encoding="utf-8"):
            pass
        return "可写，设定将存到 %s" % p

    def _end_to_end():
        """真正跑一遍改写，确认打包后还能处理文档。"""
        import tempfile
        from docx import Document
        from rewriter import rewrite
        # 造一份最小文档
        tmpd = tempfile.mkdtemp()
        src = os.path.join(tmpd, "t.docx")
        d = Document()
        d.add_paragraph("一、引言")
        d.add_paragraph("这是正文段落，用于验证改写链路。")
        d.save(src)
        out = os.path.join(tmpd, "t_out.docx")
        res = rewrite(src, out, renumber_headings=True)
        assert os.path.exists(out), "未生成输出文件"
        d2 = Document(out)
        txt = "\n".join(p.text for p in d2.paragraphs)
        assert "引言" in txt, "文字丢失"
        return "改写 %d 段，输出可重新打开" % res["changed_count"]

    def _gui_construct():
        """真正把图形界面搭起来再销毁——验证打包后的 GUI 可用。

        无图形显示的环境（CI runner、纯 SSH）建不了窗口，这属于环境限制，
        此时跳过而不是判为失败。踩过的坑：GitHub ubuntu runner 上抛
        TclError: no display name and no $DISPLAY environment variable。
        """
        try:
            from environment_probe import can_create_window
            can_gui, why = can_create_window()
        except Exception:
            can_gui, why = True, ""
        if not can_gui:
            return "跳过：当前环境无图形显示（%s）" % why
        import tkinter as tk
        import app as appmod
        root = tk.Tk()
        root.withdraw()                 # 不显示，避免弹窗
        a = appmod.App(root)
        n_region = len(a.preset)
        n_rows = len(a.vars)
        assert n_region >= 10, "预设区域过少"
        assert n_rows == n_region, "界面行数与预设不符"
        # 走一遍取值，确认控件都活着
        for vs in a.vars.values():
            for var in vs.values():
                var.get()
        root.destroy()
        return "界面构建成功：%d 个区域 × %d 个可调项" % (n_region, len(appmod.COLUMNS))

    def _superscript_rules():
        """核对上标的关键规则：引用要上标、条目序号绝不能上标。"""
        from superscript import is_reference_marker, find_markers
        assert is_reference_marker("[1]"), "[1] 应判为引用"
        assert is_reference_marker("[1,2]"), "[1,2] 应判为引用"
        assert is_reference_marker("[3-5]"), "[3-5] 应判为引用"
        assert not is_reference_marker("[2010-2015]"), "年份区间不应判为引用"
        assert not is_reference_marker("[10000]"), "五位数不应判为引用"
        got = [t for _, _, t in find_markers("见[1]与②，另见[2010-2015]。")]
        assert "[1]" in got and "②" in got, "引用与注释应被识别"
        assert "[2010-2015]" not in got, "年份区间不应被识别为标记"
        return "引用/注释识别正确，年份区间已排除"

    check("运行环境", _python)
    check("运行模式", _frozen)
    check("python-docx", _docx)
    check("lxml", _lxml)
    check("tkinter", _tk)
    check("业务模块", _modules)
    check("预设读写", _preset)
    check("用户目录可写", _writable)
    check("上标规则", _superscript_rules)
    check("端到端改写", _end_to_end)
    check("图形界面构建", _gui_construct)

    out("-" * 62)
    out("结论: %s" % ("全部通过，可以分发" if ok else "存在失败项，需修复"))
    out("=" * 62)

    # 打包成 --windowed 后没有控制台，把报告写到 exe 同目录，便于验证
    try:
        from app import user_dir
        rp = os.path.join(user_dir(), "自检报告.txt")
    except Exception:
        rp = os.path.join(os.path.dirname(os.path.abspath(sys.executable)), "自检报告.txt")
    try:
        with open(rp, "w", encoding="utf-8") as f:
            f.write("\n".join(lines) + "\n")
    except Exception:
        pass
    return 0 if ok else 1


def main():
    if "--selftest" in sys.argv:
        sys.exit(selftest())
    import app
    import tkinter as tk
    root = tk.Tk()
    app.App(root)
    root.mainloop()


if __name__ == "__main__":
    main()
